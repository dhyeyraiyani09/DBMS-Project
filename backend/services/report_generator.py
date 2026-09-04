import io
from docx import Document
from fpdf import FPDF
from backend.models import AnalysisHistory
import json

class ReportGenerator:
    def __init__(self, history: AnalysisHistory):
        self.history = history

    def _prepare_text_lines(self):
        lines = []
        lines.append("="*50)
        lines.append(f"INDEX RECOMMENDATION ENGINE - REPORT")
        lines.append(f"Analysis ID: {self.history.id}")
        lines.append(f"Date: {self.history.created_at}")
        lines.append("="*50)
        lines.append("")
        
        lines.append("1. INPUT SQL")
        lines.append("-" * 20)
        lines.append(self.history.query)
        lines.append("")

        if self.history.analysis_result:
            lines.append("2. QUERY ANALYSIS")
            lines.append("-" * 20)
            res = self.history.analysis_result
            lines.append(f"Tables: {[t.get('name') for t in res.get('tables', [])]}")
            lines.append(f"WHERE Conditions: {len(res.get('where_conditions', []))}")
            for w in res.get('where_conditions', []):
                lines.append(f"  - {w.get('column')} {w.get('operator')} (Type: {w.get('classification')})")
            lines.append("")

        if self.history.recommendation_result:
            lines.append("3. INDEX CANDIDATES")
            lines.append("-" * 20)
            res = self.history.recommendation_result
            for c in res.get("candidates", []):
                lines.append(f"Candidate: {c.get('table')}({', '.join(c.get('columns', []))})")
                lines.append(f"Score: {c.get('score')} | Status: {c.get('status')}")
                for f in c.get('factors', []):
                    sign = "+" if f.get('points', 0) > 0 else ""
                    lines.append(f"  {sign}{f.get('points')}: {f.get('reason')}")
                lines.append("")

            lines.append("4. FINAL RECOMMENDATION")
            lines.append("-" * 20)
            recs = res.get("recommendations", [])
            if recs:
                for r in recs:
                    lines.append(f"RECOMMENDED: {r.get('table')}({', '.join(r.get('columns', []))}) with score {r.get('score')}")
            else:
                lines.append("No new index recommended based on deterministic threshold.")
            lines.append("")

        lines.append("5. MYSQL EXPLAIN")
        lines.append("-" * 20)
        if self.history.explain_result and self.history.explain_result.get("success"):
            rows = self.history.explain_result.get("explain_rows", [])
            if not rows:
                lines.append("No EXPLAIN rows returned.")
            for i, r in enumerate(rows, 1):
                lines.append(f"Row {i}: Table={r.get('table')}, Type={r.get('type')}, Key={r.get('key')}, Rows={r.get('rows')}")
                lines.append(f"  Extra: {r.get('Extra')}")
                if r.get("interpretation"):
                    lines.append(f"  Interpretation: {r.get('interpretation')}")
        else:
            lines.append("EXPLAIN was not available for this analysis.")
        lines.append("")
        
        lines.append("EDUCATIONAL NOTE")
        lines.append("-" * 20)
        lines.append("Heuristic recommendation scores are distinct from the MySQL optimizer's EXPLAIN information.")
        lines.append("Actual execution timing (measured in ms) is not included unless physically benchmarked.")
        
        return lines

    def generate_txt(self) -> bytes:
        lines = self._prepare_text_lines()
        return "\n".join(lines).encode("utf-8")

    def generate_pdf(self) -> bytes:
        lines = self._prepare_text_lines()
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=10)
        for line in lines:
            # Handle non-latin-1 characters by replacing them or ignoring
            safe_line = line.encode('latin-1', 'replace').decode('latin-1')
            pdf.cell(0, 5, safe_line, new_x="LMARGIN", new_y="NEXT")
        return bytes(pdf.output())

    def generate_docx(self) -> bytes:
        lines = self._prepare_text_lines()
        doc = Document()
        doc.add_heading(f"Index Recommendation Engine Report (ID: {self.history.id})", level=1)
        
        for line in lines:
            doc.add_paragraph(line)
            
        f = io.BytesIO()
        doc.save(f)
        return f.getvalue()
